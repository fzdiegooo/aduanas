import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

directorio_base = os.path.join(os.path.dirname(__file__), "data")
os.makedirs(directorio_base, exist_ok=True)
archivo_excel = os.path.join(directorio_base, "detalles_manifiestos_aereos.xlsx")
archivo_actualizado = os.path.join(directorio_base, "detalle_aereo_completo.xlsx")

PALABRAS_CLAVE_PRODUCTOS = {
    "Jengibre fresco": ["jengibre", "ginger", "organic ginger", "fresh organic ginger", "organic fresh ginger"],
    "Cúrcuma": ["turmeric", "curcuma", "cúrcuma", "organic turmeric"],
    "Palta": ["palta", "paltas", "avocado", "aguacate", "aguacates"],
    "Arándano": ["arándano", "arándanos", "blueberries"],
    "Naranja": ["naranja", "naranjas", "orange", "oranges"],
    "Café": ["cafe", "café", "cafes", "coffee"],
    "Cacao": ["cacao", "cocoa"],
    "Uva": ["uva", "uvas", "grape", "grapes"],
    "Mango": ["mango", "mangos"],
    "Achiote": ["achiote", "annatto"],
    "Ajo": ["ajo", "garlic"],
    "Limón": ["limon", "lemon", "limón", "lemons"],
    "Esparrago": ["esparrago", "esparragos"],
    "Fresas": ["Fresa", "Strawberries", "Fresas"],
    "Quinua": ["Quinua", "Quinoa"],
    "Mandarina": ["Mandarina", "Clementinas", "Clementine"],
    "Menestras": ["Beans", "lentejas", "menestras", "frejoles", "alberja"],
    "Cebolla": ["Cebolla", "Onion", "Onions", "Cebollas"]
}

EXPORTADORES_ESPECIFICOS = [
    "vancard", "elisur", "fruitxchange", "hamillton", "la grama", "fruticola selva",
    "namaskar", "masojo", "anawi", "jungle fresh organic", "la campiña",
    "jch exportaciones", "blue pacific oils", "aseptic peruvian fruit"
]

IMPORTADORES_PAISES = {
    "AXARFRUIT": "ESPAÑA",
    "JALHUCA EXPLOTACIONES SL": "ESPAÑA",
    "PIMENTON Y DERIVADOS S.L CALLE CATA": "ESPAÑA",
    "CASCONE": "PAISES BAJOS",
    "FRUCHTHANSA": "PAISES BAJOS",
    "IPOKI BV": "PAISES BAJOS",
    "NATURE S PRODUCE SP Z": "PAISES BAJOS",
    "NFG New Fruit Group": "PAISES BAJOS",
    "VISION INTERNATIONAL B.V.": "PAISES BAJOS",
    "AGROFAIR": "ITALIA",
    "ANAWI USA": "ESTADOS UNIDOS",
    "DELINA": "ESTADOS UNIDOS",
    "ECORIPE TROPICALS": "ESTADOS UNIDOS",
    "GLOBAL FARMS ENTERPRISES": "ESTADOS UNIDOS",
    "HEATH AND LEJEUNE": "ESTADOS UNIDOS",
    "I LOVE PRODUCE": "ESTADOS UNIDOS",
    "INTERNATIONAL SPECIALTY PRODUCE": "ESTADOS UNIDOS",
    "IPOKI PRODUCE LLC": "ESTADOS UNIDOS",
    "J Y C TROPICALS INC": "ESTADOS UNIDOS",
    "JLZ PRODUCE": "ESTADOS UNIDOS",
    "SUNDINE PRODUCE INC": "ESTADOS UNIDOS",
    "TRINITY DISTRIBUTION INC": "ESTADOS UNIDOS",
    "UREN NORTH AMERICA LLC": "ESTADOS UNIDOS",
    "VIVA TIERRA ORGANIC INC": "ESTADOS UNIDOS",
    "THOMAS FRESH INC": "CANADÁ",
    "SOL FRUIT IMPORTS LTD": "REINO UNIDO"
}

CONTINENTES = {
    "ESPAÑA": "Europa", "PAISES BAJOS": "Europa", "ITALIA": "Europa",
    "ESTADOS UNIDOS": "América", "CANADÁ": "América", "REINO UNIDO": "Europa"
}

if not os.path.exists(archivo_excel):
    print(f"El archivo {archivo_excel} no existe.")
else:
    df = pd.read_excel(archivo_excel)

    def identificar_producto(descripcion, embarcador, consignatario):
        if pd.isna(descripcion) or not isinstance(descripcion, str):
            descripcion = ""
        descripcion = descripcion.lower()

        for producto, palabras_clave in PALABRAS_CLAVE_PRODUCTOS.items():
            if any(palabra in descripcion for palabra in palabras_clave):
                return producto

        if pd.notna(embarcador) and isinstance(embarcador, str):
            for exportador in EXPORTADORES_ESPECIFICOS:
                if exportador.lower() in embarcador.lower():
                    return "¿Jengibre?"

        if pd.notna(consignatario) and isinstance(consignatario, str):
            for importador in IMPORTADORES_PAISES.keys():
                if importador.lower() in consignatario.lower():
                    return "¿Jengibre?"

        return "No asociado"

    def segmentar_jengibre(descripcion, producto):
        if producto == "Jengibre fresco":
            descripcion = descripcion.lower()
            if "juice" in descripcion or "jugo" in descripcion:
                return "Jugo de jengibre"
            if "bags" in descripcion or "sacos" in descripcion or "bolsas" in descripcion:
                return "Jengibre deshidratado"
        return producto

    def identificar_pais(consignatario):
        if pd.isna(consignatario) or not isinstance(consignatario, str):
            return "No asociado"
        for importador, pais in IMPORTADORES_PAISES.items():
            if importador.lower() in consignatario.lower():
                return pais
        return "No asociado"

    def identificar_continente(pais):
        if pd.isna(pais):
            return "No asociado"
        return CONTINENTES.get(pais, "No asociado")

    df["Producto"] = df.apply(
        lambda row: segmentar_jengibre(
            row.get("Descripción de Mercadería", ""),
            identificar_producto(
                row.get("Descripción de Mercadería", ""),
                row.get("Embarcador", ""),
                row.get("Consignatario", "")
            )
        ), axis=1
    )

    df["Pais"] = df["Consignatario"].apply(identificar_pais)
    df["Continente"] = df["Pais"].apply(identificar_continente)

    df.to_excel(archivo_actualizado, index=False)

    wb = load_workbook(archivo_actualizado)
    ws = wb.active
    tabla = Table(
        displayName="DetalleAereoCompleto",
        ref=f"A1:{chr(65 + len(df.columns) - 1)}{len(df) + 1}"
    )
    estilo = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=True
    )
    tabla.tableStyleInfo = estilo
    ws.add_table(tabla)
    wb.save(archivo_actualizado)

    print(f"Archivo actualizado y convertido a tabla guardado como '{archivo_actualizado}'.")